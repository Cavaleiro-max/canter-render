# api.py
# Requisitos: pip install flask flask-socketio requests reportlab python-docx openpyxl pillow matplotlib werkzeug eventlet flask-cors
import eventlet
eventlet.monkey_patch()
import os
import re
import uuid
import time
import json
import textwrap
import mimetypes
from pathlib import Path
from io import BytesIO
from flask import Flask, request, send_file, jsonify
from flask_socketio import SocketIO, emit
from werkzeug.utils import secure_filename
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
import requests
from flask_cors import CORS

# ----------------- CONFIG -----------------
APP_DIR = Path(__file__).parent
OUTPUT_DIR = APP_DIR / "generated_files"
UPLOAD_DIR = APP_DIR / "uploads"
OUTPUT_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(exist_ok=True)

OLLAMA_URL = "http://localhost:11434/api/generate"
DEFAULT_MODEL = "openchat"

ALLOWED_EXT = {"png", "jpg", "jpeg", "pdf", "docx", "xlsx", "txt", "csv", "mp4"}
ALLOWED_MIME = {
    "image/png", "image/jpeg", "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/plain", "text/csv", "video/mp4"
}

DEFAULT_FONT_PATHS = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
]
FONT_PATH = next((p for p in DEFAULT_FONT_PATHS if Path(p).exists()), None)

# ----------------- MODELO INTELIGENTE -----------------
MODEL_ROUTING = {
    "openchat": "openchat",
    "llama": "llama3:8b"
}

def escolher_modelo(prompt: str) -> str:
    prompt_lower = prompt.lower()
    if any(x in prompt_lower for x in [
        "código", "programa", "função", "algoritmo", "explica isso", "analisa",
        "resuma", "documento", "tabela", "planilha", "dados", "estrutura", "modelo"
    ]):
        return MODEL_ROUTING["llama"]
    if any(x in prompt_lower for x in [
        "conversa", "emoção", "sentimento", "diálogo", "chat", "resposta longa",
        "explicação pessoal", "história", "personagem", "simule", "interaja"
    ]):
        return MODEL_ROUTING["openchat"]
    return DEFAULT_MODEL

# ----------------- FLASK & SOCKET -----------------
app = Flask(__name__)
CORS(app, origins="*")
socketio = SocketIO(app, cors_allowed_origins="*", async_mode="eventlet")

# ----------------- HELPERS -----------------
def _unique_name(prefix: str, ext: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8]}.{ext}"

def _emit_progress_to_sid(sid: str, pct: int = 0, msg: str = ""):
    socketio.emit('progress', {"progress": int(pct), "message": msg}, room=sid)

def _safe_ext_from_filename(filename: str) -> str:
    return secure_filename(filename).rsplit('.', 1)[-1].lower() if '.' in filename else ''

def _validate_upload(file_storage) -> tuple[bool, str]:
    content_type = file_storage.mimetype or mimetypes.guess_type(file_storage.filename)[0] or ""
    ext = _safe_ext_from_filename(file_storage.filename)
    if ext not in ALLOWED_EXT:
        return False, f"Extensão '{ext}' não permitida"
    if content_type and content_type not in ALLOWED_MIME:
        return False, f"MIME '{content_type}' não permitido"
    return True, ""

# ----------------- GERAÇÃO DE ARQUIVOS -----------------
def gerar_pdf(texto: str, nome: str | None = None, sid: str | None = None) -> str:
    nome = nome or _unique_name("document", "pdf")
    caminho = OUTPUT_DIR / nome
    c = canvas.Canvas(str(caminho))
    margin_left, margin_top, line_height, y = 50, 800, 14, 800
    lines = []
    for p in texto.splitlines():
        wrapped = textwrap.wrap(p, width=110)
        lines.extend(wrapped or [""])
    total = max(1, len(lines))
    for idx, line in enumerate(lines, start=1):
        c.drawString(margin_left, y, line)
        y -= line_height
        if y < 50:
            c.showPage()
            y = margin_top
        if sid:
            _emit_progress_to_sid(sid, int((idx / total) * 85), f"Escrevendo PDF {idx}/{total}")
    c.save()
    if sid:
        _emit_progress_to_sid(sid, 100, "PDF pronto")
    return str(caminho)

def gerar_docx(texto: str, nome: str | None = None, sid: str | None = None) -> str:
    nome = nome or _unique_name("document", "docx")
    caminho = OUTPUT_DIR / nome
    doc = Document()
    lines = texto.splitlines() or [texto]
    total = max(1, len(lines))
    for idx, line in enumerate(lines, start=1):
        doc.add_paragraph(line)
        if sid:
            _emit_progress_to_sid(sid, int((idx / total) * 90), f"Gerando DOCX {idx}/{total}")
    doc.save(str(caminho))
    if sid:
        _emit_progress_to_sid(sid, 100, "DOCX pronto")
    return str(caminho)

def gerar_xlsx(texto: str, nome: str | None = None, sid: str | None = None) -> str:
    nome = nome or _unique_name("sheet", "xlsx")
    caminho = OUTPUT_DIR / nome
    wb = Workbook()
    ws = wb.active
    lines = texto.splitlines() or [texto]
    total = max(1, len(lines))
    for i, line in enumerate(lines, start=1):
        cols = re.split(r'\t|\s{2,}|;', line)
        for j, cell in enumerate(cols, start=1):
            ws.cell(row=i, column=j, value=cell.strip())
        if sid:
            _emit_progress_to_sid(sid, int((i / total) * 90), f"Escrevendo planilha {i}/{total}")
    wb.save(str(caminho))
    if sid:
        _emit_progress_to_sid(sid, 100, "Planilha pronta")
    return str(caminho)

def gerar_png_from_text(texto: str, nome: str | None = None, sid: str | None = None) -> str:
    nome = nome or _unique_name("image", "png")
    caminho = OUTPUT_DIR / nome
    width, height, margin, line_spacing = 1200, 1600, 40, 6
    bg = (255, 255, 255)
    font = ImageFont.truetype(FONT_PATH, size=20) if FONT_PATH else ImageFont.load_default()
    paragraphs = texto.splitlines()
    img = Image.new("RGB", (width, height), color=bg)
    draw = ImageDraw.Draw(img)
    y = margin
    total_lines = sum(len(textwrap.wrap(p, width=80)) or 1 for p in paragraphs)
    drawn = 0
    for para in paragraphs:
        wrapped = textwrap.wrap(para, width=80) or [""]
        for line in wrapped:
            draw.text((margin, y), line, fill=(20, 20, 20), font=font)
            y += font.getsize(line)[1] + line_spacing
            drawn += 1
            if sid:
                _emit_progress_to_sid(sid, int((drawn / total_lines) * 90), f"Gerando imagem {drawn}/{total_lines}")
            if y > height - margin:
                new_img = Image.new("RGB", (width, img.height + height), color=bg)
                new_img.paste(img, (0, 0))
                img = new_img
                draw = ImageDraw.Draw(img)
                y = margin
        y += 8
    img.save(str(caminho), format="PNG", optimize=True)
    if sid:
        _emit_progress_to_sid(sid, 100, "Imagem pronta")
    return str(caminho)

# ----------------- CHAMADA AO OLLAMA -----------------
def call_ollama_stream(prompt: str, sid: str | None = None, retries: int = 3, backoff: float = 1.0) -> str:
    modelo_usado = escolher_modelo(prompt)
    payload = {"model": modelo_usado, "prompt": prompt, "stream": True}
    attempt = 0
    text_acc = ""
    while attempt < retries:
        try:
            r = requests.post(OLLAMA_URL, json=payload, stream=True, timeout=300)
            r.raise_for_status()
            for chunk in r.iter_lines():
                if not chunk:
                    continue
                try:
                    decoded = chunk.decode()
                except Exception:
                    continue
                try:
                    j = json.loads(decoded)
                    for k in ("response", "text", "content"):
                        v = j.get(k)
                        if isinstance(v, str):
                            text_acc += v
                            if sid:
                                socketio.emit('response', v, room=sid)
                            break
                    else:
                        text_acc += decoded
                        if sid:
                            socketio.emit('response', decoded, room=sid)
                except Exception:
                    text_acc += decoded
                    if sid:
                        socketio.emit('response', decoded, room=sid)
            return text_acc
        except Exception as e:
            attempt += 1
            time.sleep(backoff * attempt)
            if attempt >= retries:
                if sid:
                    socketio.emit('response', f"❌ Falha ao comunicar com Ollama: {e}", room=sid)
                raise
    return text_acc

# ----------------- SOCKET & HTTP HANDLERS -----------------
@socketio.on('connect')
def on_connect():
    sid = request.sid
    emit('response', '🧠 Servidor pronto. Envie texto ou faça upload em /upload.', room=sid)

@socketio.on('message')
def handle_message(data):
    sid = request.sid
    prompt = data.get("command") if isinstance(data, dict) else (data or "")
    if not prompt:
        socketio.emit('response', '⚠️ Nenhum comando recebido.', room=sid)
        return
    try:
        _emit_progress_to_sid(sid, pct=1, msg="Chamando modelo...")
        full_text = call_ollama_stream(prompt, sid=sid)
        _emit_progress_to_sid(sid, pct=70, msg="Modelo respondeu, preparando saída...")
        lower = prompt.lower()
        if any(k in lower for k in ("gerar pdf", "gere um pdf", "gerar um pdf")):
            path = gerar_pdf(full_text, sid=sid)
            filename = os.path.basename(path)
            socketio.emit('file', {"filename": filename, "url": f"/download/{filename}"}, room=sid)
            socketio.emit('response', f"✅ PDF gerado: {filename}", room=sid)
            return
        if any(k in lower for k in ("gerar docx", "gere um docx", "gerar um doc")):
            path = gerar_docx(full_text, sid=sid)
            filename = os.path.basename(path)
            socketio.emit('file', {"filename": filename, "url": f"/download/{filename}"}, room=sid)
            socketio.emit('response', f"✅ DOCX gerado: {filename}", room=sid)
            return
        if any(k in lower for k in ("planilha", "gerar xlsx", "gerar planilha")):
            path = gerar_xlsx(full_text, sid=sid)
            filename = os.path.basename(path)
            socketio.emit('file', {"filename": filename, "url": f"/download/{filename}"}, room=sid)
            socketio.emit('response', f"✅ XLSX gerado: {filename}", room=sid)
            return
        if any(k in lower for k in ("gerar imagem", "gerar png", "imagem")):
            path = gerar_png_from_text(full_text, sid=sid)
            filename = os.path.basename(path)
            socketio.emit('file', {"filename": filename, "url": f"/download/{filename}"}, room=sid)
            socketio.emit('response', f"🖼️ Imagem gerada: {filename}", room=sid)
            return
        socketio.emit('response', full_text, room=sid)
    except Exception as e:
        socketio.emit('response', f"❌ Erro interno: {e}", room=sid)

@app.route('/upload', methods=['POST'])
def upload_file():
    sid = request.form.get('sid') or request.args.get('sid')
    if 'file' not in request.files:
        return jsonify({"error": "no file"}), 400
    f = request.files['file']
    ok, msg = _validate_upload(f)
    if not ok:
        return jsonify({"error": msg}), 400
    filename = secure_filename(f.filename)
    unique = f"{uuid.uuid4().hex[:8]}-{filename}"
    save_path = UPLOAD_DIR / unique
    f.save(save_path)
    ext = _safe_ext_from_filename(unique)
    if ext in {"png", "jpg", "jpeg"}:
        try:
            img = Image.open(save_path)
            img.thumbnail((1600, 1600))
            thumb_name = f"thumb-{unique}"
            thumb_path = OUTPUT_DIR / thumb_name
            img.save(thumb_path)
            if sid:
                socketio.emit('uploaded', {"filename": unique, "url": f"/download/{thumb_name}"}, room=sid)
            else:
                socketio.emit('uploaded', {"filename": unique, "url": f"/download/{thumb_name}"})
        except Exception:
            if sid:
                socketio.emit('uploaded', {"filename": unique, "url": f"/download/{unique}"}, room=sid)
            else:
                socketio.emit('uploaded', {"filename": unique, "url": f"/download/{unique}"})
    else:
        if sid:
            socketio.emit('uploaded', {"filename": unique, "url": f"/download/{unique}"}, room=sid)
        else:
            socketio.emit('uploaded', {"filename": unique, "url": f"/download/{unique}"})
    return jsonify({"ok": True, "filename": unique}), 200

@app.route('/download/<path:filename>')
def download_file(filename):
    p1 = OUTPUT_DIR / filename
    p2 = UPLOAD_DIR / filename
    if p1.exists():
        return send_file(str(p1), as_attachment=True)
    if p2.exists():
        return send_file(str(p2), as_attachment=True)
    return "Arquivo não encontrado", 404

@app.route('/list_generated')
def list_generated():
    files = [f.name for f in OUTPUT_DIR.iterdir() if f.is_file()]
    return jsonify(files)

@app.route('/list_uploads')
def list_uploads():
    files = [f.name for f in UPLOAD_DIR.iterdir() if f.is_file()]
    return jsonify(files)

if __name__ == "__main__":
    socketio.run(app, host="0.0.0.0", port=5000)
